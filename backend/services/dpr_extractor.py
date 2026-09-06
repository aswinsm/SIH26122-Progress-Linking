import csv
import os
from pathlib import Path
from typing import Literal

from dotenv import load_dotenv
from groq import Groq
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pypdf import PdfReader


# ============================================================
# PROJECT / ENVIRONMENT SETUP
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]

ENV_FILE = PROJECT_ROOT / ".env"

load_dotenv(ENV_FILE)

GROQ_API_KEY = os.getenv("GROQ_API_KEY")

GROQ_MODEL = os.getenv(
    "GROQ_MODEL",
    "openai/gpt-oss-20b"
)

if not GROQ_API_KEY:
    raise ValueError(
        "GROQ_API_KEY was not found in the project .env file"
    )

client = Groq(
    api_key=GROQ_API_KEY
)


# ============================================================
# DATA MODELS
# ============================================================

DPRStatus = Literal[
    "Completed",
    "Started",
    "In Progress",
    "Delayed",
    "Not Started"
]


class DPRData(BaseModel):

    activity_description: str | None = None
    discipline: str | None = None
    status: DPRStatus | None = None
    tag: str | None = None
    date: str | None = None


class DPRExtractionResult(BaseModel):

    activities: list[DPRData] = Field(
        default_factory=list
    )


# ============================================================
# GROQ STRUCTURED OUTPUT SCHEMA
# ============================================================

DPR_RESPONSE_SCHEMA = {

    "type": "object",

    "properties": {

        "activities": {

            "type": "array",

            "items": {

                "type": "object",

                "properties": {

                    "activity_description": {
                        "type": [
                            "string",
                            "null"
                        ]
                    },

                    "discipline": {
                        "type": [
                            "string",
                            "null"
                        ]
                    },

                    "status": {

                        "anyOf": [

                            {
                                "type": "string",
                                "enum": [
                                    "Completed",
                                    "Started",
                                    "In Progress",
                                    "Delayed",
                                    "Not Started"
                                ]
                            },

                            {
                                "type": "null"
                            }
                        ]
                    },

                    "tag": {
                        "type": [
                            "string",
                            "null"
                        ]
                    },

                    "date": {
                        "type": [
                            "string",
                            "null"
                        ]
                    }
                },

                "required": [
                    "activity_description",
                    "discipline",
                    "status",
                    "tag",
                    "date"
                ],

                "additionalProperties": False
            }
        }
    },

    "required": [
        "activities"
    ],

    "additionalProperties": False
}


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_text_from_pdf(
    pdf_path: str | Path
) -> str:

    pdf_path = Path(pdf_path)

    if not pdf_path.exists():
        raise FileNotFoundError(
            f"PDF file not found: {pdf_path}"
        )

    reader = PdfReader(
        pdf_path
    )

    text_parts = []

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):

        page_text = page.extract_text()

        if page_text:

            text_parts.append(
                f"\n--- Page {page_number} ---\n"
                f"{page_text}"
            )

    return "\n".join(
        text_parts
    )


# ============================================================
# EXCEL TEXT EXTRACTION
# ============================================================

def extract_text_from_excel(
    excel_path: str | Path
) -> str:

    excel_path = Path(excel_path)

    if not excel_path.exists():
        raise FileNotFoundError(
            f"Excel file not found: {excel_path}"
        )

    workbook = load_workbook(
        excel_path,
        data_only=True,
        read_only=True
    )

    text_parts = []

    try:

        for sheet in workbook.worksheets:

            text_parts.append(
                f"\n--- Sheet: {sheet.title} ---"
            )

            for row in sheet.iter_rows(
                values_only=True
            ):

                values = []

                for cell in row:

                    if cell is None:
                        continue

                    value = str(cell).strip()

                    if value:
                        values.append(value)

                if values:

                    text_parts.append(
                        " | ".join(values)
                    )

    finally:

        workbook.close()

    return "\n".join(
        text_parts
    )


# ============================================================
# CSV TEXT EXTRACTION
# ============================================================

def extract_text_from_csv(
    csv_path: str | Path
) -> str:

    csv_path = Path(csv_path)

    if not csv_path.exists():
        raise FileNotFoundError(
            f"CSV file not found: {csv_path}"
        )

    text_parts = []

    with csv_path.open(
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline=""
    ) as file:

        reader = csv.reader(file)

        for row in reader:

            values = []

            for cell in row:

                value = str(cell).strip()

                if value:
                    values.append(value)

            if values:

                text_parts.append(
                    " | ".join(values)
                )

    return "\n".join(
        text_parts
    )


# ============================================================
# AI DPR EXTRACTION
# ============================================================

def extract_dpr(
    dpr_text: str,
    reference_date: str | None = None
) -> DPRExtractionResult:

    if (
        not dpr_text
        or not dpr_text.strip()
    ):

        raise ValueError(
            "DPR text cannot be empty"
        )

    # --------------------------------------------------------
    # DATE INSTRUCTION
    # --------------------------------------------------------

    if reference_date:

        reference_date_instruction = f"""
Reference date: {reference_date}

If expressions such as "today" or "yesterday"
appear, interpret them relative to this date.
"""

    else:

        reference_date_instruction = """
No reference date is available.

If only a relative date such as "today" or
"yesterday" appears and the actual date cannot
be determined, return null.
"""

    # --------------------------------------------------------
    # PROMPT
    # --------------------------------------------------------

    prompt = f"""
You are an engineering Daily Progress Report
(DPR) extraction system for infrastructure
construction projects.

Extract construction progress activities from
the supplied DPR content.

The report may contain Civil, Piping,
Mechanical, Electrical, Instrumentation,
HSE or other construction activities.

IMPORTANT RULES:

1. A DPR may contain MULTIPLE activities.
2. Create one separate object for every distinct
   construction progress activity.
3. Never invent information.
4. Return null when information is unavailable.
5. Ignore headings, administrative text and
   unrelated spreadsheet information.

For every activity identify:

- activity_description
- discipline
- status
- tag
- date


ACTIVITY DESCRIPTION:

Keep it short and meaningful.

Examples:

"F101 foundation concreting"
"Reinforcement for F102"
"Cable tray installation Area A"


DISCIPLINE:

Examples include:

Civil
Piping
Mechanical
Electrical
Instrumentation
HSE

Only infer discipline when engineering context
strongly supports it.


STATUS:

Use ONLY these values:

Completed
Started
In Progress
Delayed
Not Started

Examples:

finished / completed / work done
→ Completed

started / commenced / began
→ Started

ongoing / underway / in progress
→ In Progress

delayed / held up / stopped due to
→ Delayed

not started / yet to start
→ Not Started


TAG:

Preserve identifiers such as:

F101
F102
L24
P-104
ST-03
Area-A

Do not invent a tag.


DATE:

Convert explicit dates to YYYY-MM-DD.

Never invent a date.

{reference_date_instruction}


DPR CONTENT:

---------------- START ----------------

{dpr_text}

---------------- END ----------------
"""

    # --------------------------------------------------------
    # GROQ REQUEST
    # --------------------------------------------------------

    try:

        response = client.chat.completions.create(

            model=GROQ_MODEL,

            messages=[
                {
                    "role": "system",
                    "content": (
                        "Extract structured engineering DPR "
                        "activities and follow the supplied "
                        "JSON schema exactly."
                    )
                },
                {
                    "role": "user",
                    "content": prompt
                }
            ],

            response_format={
                "type": "json_schema",
                "json_schema": {
                    "name": "dpr_extraction",
                    "strict": True,
                    "schema": DPR_RESPONSE_SCHEMA
                }
            },

            temperature=0
        )

    except Exception as exc:

        raise RuntimeError(
            f"Groq DPR extraction failed: {exc}"
        ) from exc

    # --------------------------------------------------------
    # READ RESPONSE
    # --------------------------------------------------------

    result = (
        response
        .choices[0]
        .message
        .content
    )

    if not result:

        raise RuntimeError(
            "Groq returned an empty response"
        )

    # --------------------------------------------------------
    # PYDANTIC VALIDATION
    # --------------------------------------------------------

    try:

        return (
            DPRExtractionResult
            .model_validate_json(
                result
            )
        )

    except Exception as exc:

        raise RuntimeError(
            "Groq response could not be "
            "validated as DPR data. "
            f"Raw response: {result}"
        ) from exc


# ============================================================
# SAFE CHUNK EXTRACTION
# ============================================================

def extract_chunk_safely(
    lines: list[str],
    sheet_name: str,
    reference_date: str | None = None
) -> list[DPRData]:

    if not lines:
        return []

    chunk_text = (
        f"Excel Sheet: {sheet_name}\n\n"
        + "\n".join(lines)
    )

    try:

        result = extract_dpr(
            chunk_text,
            reference_date=reference_date
        )

        return result.activities

    except Exception:

        # If the chunk is already one row,
        # retrying by splitting further is impossible.
        if len(lines) == 1:
            raise

        # Split failed chunk into two smaller pieces.
        middle = len(lines) // 2

        first_half = lines[:middle]
        second_half = lines[middle:]

        activities = []

        activities.extend(
            extract_chunk_safely(
                first_half,
                sheet_name,
                reference_date
            )
        )

        activities.extend(
            extract_chunk_safely(
                second_half,
                sheet_name,
                reference_date
            )
        )

        return activities


# ============================================================
# EXCEL DPR CHUNK PROCESSING
# ============================================================

def extract_dpr_from_excel_chunks(
    excel_path: str | Path,
    reference_date: str | None = None,
    chunk_size: int = 8
) -> DPRExtractionResult:

    excel_path = Path(excel_path)

    if not excel_path.exists():

        raise FileNotFoundError(
            f"Excel file not found: {excel_path}"
        )

    workbook = load_workbook(
        excel_path,
        data_only=True,
        read_only=True
    )

    all_activities = []

    try:

        for sheet in workbook.worksheets:

            rows = []

            # ------------------------------------------------
            # CONVERT EXCEL ROWS TO TEXT
            # ------------------------------------------------

            for row in sheet.iter_rows(
                values_only=True
            ):

                values = []

                for cell in row:

                    if cell is None:
                        continue

                    value = str(cell).strip()

                    if value:

                        values.append(
                            value
                        )

                if values:

                    rows.append(
                        " | ".join(values)
                    )

            # ------------------------------------------------
            # IGNORE EMPTY SHEETS
            # ------------------------------------------------

            if not rows:
                continue

            # ------------------------------------------------
            # PROCESS SMALL CHUNKS
            # ------------------------------------------------

            for start in range(
                0,
                len(rows),
                chunk_size
            ):

                chunk = rows[
                    start:
                    start + chunk_size
                ]

                activities = (
                    extract_chunk_safely(
                        lines=chunk,
                        sheet_name=sheet.title,
                        reference_date=reference_date
                    )
                )

                all_activities.extend(
                    activities
                )

    finally:

        workbook.close()

    # --------------------------------------------------------
    # REMOVE DUPLICATES
    # --------------------------------------------------------

    unique_activities = []

    seen = set()

    for activity in all_activities:

        key = (

            (
                activity.activity_description
                or ""
            ).lower().strip(),

            (
                activity.discipline
                or ""
            ).lower().strip(),

            (
                activity.status
                or ""
            ).lower().strip(),

            (
                activity.tag
                or ""
            ).lower().strip(),

            (
                activity.date
                or ""
            ).lower().strip()
        )

        if key not in seen:

            seen.add(key)

            unique_activities.append(
                activity
            )

    return DPRExtractionResult(
        activities=unique_activities
    )


# ============================================================
# FILE DPR EXTRACTION
# ============================================================

def extract_dpr_from_file(
    file_path: str | Path,
    reference_date: str | None = None
) -> DPRExtractionResult:

    file_path = Path(file_path)

    if not file_path.exists():

        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = (
        file_path
        .suffix
        .lower()
    )

    # --------------------------------------------------------
    # EXCEL
    # --------------------------------------------------------

    if extension in [
        ".xlsx",
        ".xlsm"
    ]:

        # IMPORTANT:
        # Excel is processed in small chunks instead
        # of sending the entire workbook to Groq.

        return extract_dpr_from_excel_chunks(
            file_path,
            reference_date=reference_date
        )

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    elif extension == ".pdf":

        text = extract_text_from_pdf(
            file_path
        )

    # --------------------------------------------------------
    # CSV
    # --------------------------------------------------------

    elif extension == ".csv":

        text = extract_text_from_csv(
            file_path
        )

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    elif extension == ".txt":

        text = file_path.read_text(
            encoding="utf-8",
            errors="replace"
        )

    # --------------------------------------------------------
    # UNSUPPORTED FILE
    # --------------------------------------------------------

    else:

        raise ValueError(
            "Unsupported file type: "
            f"{extension}. "
            "Supported formats are "
            "PDF, XLSX, XLSM, CSV and TXT."
        )

    # --------------------------------------------------------
    # EMPTY CONTENT CHECK
    # --------------------------------------------------------

    if not text.strip():

        raise ValueError(
            "No readable text was found "
            "inside the uploaded file."
        )

    # --------------------------------------------------------
    # NORMAL TEXT EXTRACTION
    # --------------------------------------------------------

    return extract_dpr(
        text,
        reference_date=reference_date
    )


# ============================================================
# LOCAL TEST
# ============================================================

if __name__ == "__main__":

    sample_dpr = """
    F101 concreting completed.
    Reinforcement for F102 started.
    Cable tray installation in Area A is in progress.
    """

    result = extract_dpr(
        sample_dpr
    )

    print(
        "\n========== DPR EXTRACTION RESULT ==========\n"
    )

    print(
        result.model_dump_json(
            indent=2
        )
    )